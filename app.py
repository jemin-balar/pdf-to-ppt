from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
import requests
import os
import uuid
from pdf2image import convert_from_path
from pptx import Presentation
from pptx.util import Inches
import threading
import time
from urllib.parse import urlparse

app = Flask(__name__)
CORS(app)

UPLOAD_FOLDER = 'temp_files'
MAX_FILE_SIZE = 100 * 1024 * 1024
CLEANUP_INTERVAL = 3600

os.makedirs(UPLOAD_FOLDER, exist_ok=True)

POPPLER_PATH = os.getenv("POPPLER_PATH", "/opt/homebrew/bin")
print(f"Using Poppler path: {POPPLER_PATH}")
print(f"Poppler path exists: {os.path.exists(POPPLER_PATH)}")

class PDFtoPPTConverter:
    def __init__(self):
        self.conversion_jobs = {}
    
    def download_pdf(self, pdf_url, filename):
        headers = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'}
        response = requests.get(pdf_url, headers=headers, timeout=60, stream=True)
        response.raise_for_status()
        
        total_size = 0
        with open(filename, 'wb') as f:
            for chunk in response.iter_content(chunk_size=8192):
                if chunk:
                    f.write(chunk)
                    total_size += len(chunk)
                    if total_size > MAX_FILE_SIZE:
                        raise ValueError(f'File too large. Maximum size is {MAX_FILE_SIZE/1024/1024}MB')
        
        if total_size == 0:
            raise ValueError('Downloaded file is empty')
        
        return True
    
    def convert_pdf_to_ppt(self, pdf_path, output_path):
        try:
            pages = convert_from_path(
                pdf_path,
                dpi=200,
                fmt='PNG',
                thread_count=4,
                poppler_path=POPPLER_PATH
            )
            prs = Presentation()
            
            for i, page in enumerate(pages):
                slide_layout = prs.slide_layouts[6]
                slide = prs.slides.add_slide(slide_layout)
                
                temp_img_path = f'{pdf_path}_page_{i}.png'
                page.save(temp_img_path, 'PNG', optimize=True, quality=90)
                
                # Get image dimensions to determine orientation
                from PIL import Image
                with Image.open(temp_img_path) as img:
                    img_width, img_height = img.size
                    
                # Calculate aspect ratio
                aspect_ratio = img_width / img_height
                
                # Determine if image is portrait or landscape
                is_portrait = img_height > img_width
                
                # Add the image to fill the entire slide
                slide.shapes.add_picture(temp_img_path, Inches(0), Inches(0), 
                                       slide_width, slide_height)
                
                try:
                    os.remove(temp_img_path)
                except:
                    pass
            
            prs.save(output_path)
            return True
            
        except Exception as e:
            raise Exception(f'Conversion failed: {str(e)}')
    
    def convert_async(self, job_id, pdf_url):
        try:
            self.conversion_jobs[job_id]['status'] = 'downloading'
            
            pdf_filename = os.path.join(UPLOAD_FOLDER, f'{job_id}.pdf')
            ppt_filename = os.path.join(UPLOAD_FOLDER, f'{job_id}.pptx')
            
            self.download_pdf(pdf_url, pdf_filename)
            
            self.conversion_jobs[job_id]['status'] = 'converting'
            self.convert_pdf_to_ppt(pdf_filename, ppt_filename)
            
            self.conversion_jobs[job_id]['status'] = 'completed'
            self.conversion_jobs[job_id]['ppt_file'] = ppt_filename
            self.conversion_jobs[job_id]['download_url'] = f'/download/{job_id}'
            
        except Exception as e:
            self.conversion_jobs[job_id]['status'] = 'failed'
            self.conversion_jobs[job_id]['error'] = str(e)
        finally:
            try:
                if os.path.exists(pdf_filename):
                    os.remove(pdf_filename)
            except:
                pass

converter = PDFtoPPTConverter()

def convert_local_file(job_id, pdf_path, ppt_path):
    try:
        converter.convert_pdf_to_ppt(pdf_path, ppt_path)
        converter.conversion_jobs[job_id]['status'] = 'completed'
        converter.conversion_jobs[job_id]['ppt_file'] = ppt_path
        converter.conversion_jobs[job_id]['download_url'] = f'/download/{job_id}'
    except Exception as e:
        converter.conversion_jobs[job_id]['status'] = 'failed'
        converter.conversion_jobs[job_id]['error'] = str(e)
    finally:
        try:
            os.remove(pdf_path)
        except:
            pass

@app.route('/')
def index():
    return send_file('templates/index.html')

@app.route('/convert', methods=['POST'])
def convert_pdf():
    try:
        data = request.get_json()
        if not data or 'pdf_url' not in data:
            return jsonify({'error': 'PDF URL is required'}), 400
        
        pdf_url = data['pdf_url']
        
        parsed_url = urlparse(pdf_url)
        if not parsed_url.scheme or not parsed_url.netloc:
            return jsonify({'error': 'Invalid URL format'}), 400
        
        job_id = str(uuid.uuid4())
        converter.conversion_jobs[job_id] = {
            'status': 'queued',
            'created_at': time.time(),
            'pdf_url': pdf_url
        }
        
        thread = threading.Thread(target=converter.convert_async, args=(job_id, pdf_url))
        thread.daemon = True
        thread.start()
        
        return jsonify({
            'success': True,
            'job_id': job_id,
            'status_url': f'/status/{job_id}'
        })
    
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/upload-convert', methods=['POST'])
def upload_and_convert():
    try:
        if 'file' not in request.files:
            return jsonify({'error': 'No file provided'}), 400
        
        file = request.files['file']
        if file.filename == '':
            return jsonify({'error': 'No file selected'}), 400
        
        if not file.filename.lower().endswith('.pdf'):
            return jsonify({'error': 'Only PDF files allowed'}), 400
        
        if file.content_length and file.content_length > MAX_FILE_SIZE:
            return jsonify({'error': f'File too large. Maximum size is {MAX_FILE_SIZE/1024/1024}MB'}), 400
        
        job_id = str(uuid.uuid4())
        pdf_filename = os.path.join(UPLOAD_FOLDER, f'{job_id}.pdf')
        ppt_filename = os.path.join(UPLOAD_FOLDER, f'{job_id}.pptx')
        
        file.save(pdf_filename)
        
        converter.conversion_jobs[job_id] = {
            'status': 'converting',
            'created_at': time.time()
        }
        
        thread = threading.Thread(target=convert_local_file, args=(job_id, pdf_filename, ppt_filename))
        thread.daemon = True
        thread.start()
        
        return jsonify({
            'success': True,
            'job_id': job_id,
            'status_url': f'/status/{job_id}'
        })
    
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/convert-local', methods=['POST'])
def convert_local_path():
    try:
        data = request.get_json()
        if not data or 'pdf_path' not in data:
            return jsonify({'error': 'PDF file path is required'}), 400
        
        pdf_path = data['pdf_path']
        
        if not os.path.exists(pdf_path):
            return jsonify({'error': 'PDF file not found'}), 404
        
        if not pdf_path.lower().endswith('.pdf'):
            return jsonify({'error': 'File must be a PDF'}), 400
        
        file_size = os.path.getsize(pdf_path)
        if file_size > MAX_FILE_SIZE:
            return jsonify({'error': f'File too large. Maximum size is {MAX_FILE_SIZE/1024/1024}MB'}), 400
        
        job_id = str(uuid.uuid4())
        ppt_filename = os.path.join(UPLOAD_FOLDER, f'{job_id}.pptx')
        
        converter.conversion_jobs[job_id] = {
            'status': 'converting',
            'created_at': time.time()
        }
        
        thread = threading.Thread(target=convert_local_file, args=(job_id, pdf_path, ppt_filename))
        thread.daemon = True
        thread.start()
        
        return jsonify({
            'success': True,
            'job_id': job_id,
            'status_url': f'/status/{job_id}'
        })
    
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/status/<job_id>', methods=['GET'])
def get_status(job_id):
    if job_id not in converter.conversion_jobs:
        return jsonify({'error': 'Job not found'}), 404
    
    job = converter.conversion_jobs[job_id]
    response_data = {
        'job_id': job_id,
        'status': job['status']
    }
    
    if job['status'] == 'completed':
        response_data['download_url'] = job['download_url']
    elif job['status'] == 'failed':
        response_data['error'] = job.get('error', 'Unknown error')
    
    return jsonify(response_data)

@app.route('/download/<job_id>', methods=['GET'])
def download_ppt(job_id):
    if job_id not in converter.conversion_jobs:
        return jsonify({'error': 'Job not found'}), 404
    
    job = converter.conversion_jobs[job_id]
    if job['status'] != 'completed':
        return jsonify({'error': 'Conversion not completed yet'}), 400
    
    ppt_file = job.get('ppt_file')
    if not ppt_file or not os.path.exists(ppt_file):
        return jsonify({'error': 'File not found'}), 404
    
    return send_file(ppt_file, 
                     as_attachment=True, 
                     download_name=f'converted_{job_id}.pptx',
                     mimetype='application/vnd.openxmlformats-officedocument.presentationml.presentation')

def cleanup_old_files():
    while True:
        try:
            current_time = time.time()
            jobs_to_remove = []
            
            for job_id, job in converter.conversion_jobs.items():
                if current_time - job['created_at'] > CLEANUP_INTERVAL:
                    ppt_file = job.get('ppt_file')
                    if ppt_file and os.path.exists(ppt_file):
                        try:
                            os.remove(ppt_file)
                        except:
                            pass
                    jobs_to_remove.append(job_id)
            
            for job_id in jobs_to_remove:
                del converter.conversion_jobs[job_id]
            
            for filename in os.listdir(UPLOAD_FOLDER):
                filepath = os.path.join(UPLOAD_FOLDER, filename)
                if os.path.isfile(filepath):
                    file_age = current_time - os.path.getctime(filepath)
                    if file_age > CLEANUP_INTERVAL:
                        try:
                            os.remove(filepath)
                        except:
                            pass
        
        except Exception:
            pass
        
        time.sleep(300)

cleanup_thread = threading.Thread(target=cleanup_old_files)
cleanup_thread.daemon = True
cleanup_thread.start()

@app.route('/health', methods=['GET'])
def health_check():
    return jsonify({'status': 'healthy', 'active_jobs': len(converter.conversion_jobs)})

if __name__ == '__main__':
    app.run(debug=False, host='0.0.0.0', port=5001)
